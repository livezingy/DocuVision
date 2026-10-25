"""DOCX render mixin for the export service.

Split from the former export_service.py monolith (S1 batch 10):
``to_docx`` (python-docx lazy import, title/info/structure/text/tables/
keywords/footer); composes into ``ExportService`` whose ``__init__`` owns
``self.output_dir``.
"""

from __future__ import annotations

import os
from typing import Any, Dict

from loguru import logger

from app.services.export_table_format import format_table_export_title


class DocxRenderMixin:
    """Word (DOCX) document renderer."""

    async def to_docx(self, result: Dict[str, Any], task_id: str) -> str:
        """
        Export result to Word document
        
        Args:
            result: Processing result dictionary
            task_id: Task identifier
        
        Returns:
            Path to exported DOCX file
        """
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.style import WD_STYLE_TYPE
        except ImportError:
            logger.error("python-docx not installed, cannot export Word document")
            raise RuntimeError("python-docx not installed")
        
        output_path = os.path.join(self.output_dir, f"{task_id}_result.docx")
        
        doc = Document()
        
        doc_info = result.get('document_info', {})
        text_blocks = result.get('text_blocks', [])
        tables = result.get('tables', [])
        keywords = result.get('keywords', [])
        layout = result.get('layout', {})
        
        # Title
        title = doc.add_heading(doc_info.get('file_name', 'Document Analysis Result'), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Document Info Section
        doc.add_heading('Document Information', level=1)
        
        info_table = doc.add_table(rows=4, cols=2)
        info_table.style = 'Table Grid'
        
        info_data = [
            ('File Name', doc_info.get('file_name', 'N/A')),
            ('Processed At', doc_info.get('processed_at', 'N/A')),
            ('Total Pages', str(doc_info.get('pages', 'N/A'))),
            ('OCR Engine', result.get('engine_used', 'N/A'))
        ]
        
        for i, (label, value) in enumerate(info_data):
            info_table.rows[i].cells[0].text = label
            info_table.rows[i].cells[1].text = value
        
        doc.add_paragraph()
        
        # Document Structure Summary
        if layout:
            doc.add_heading('Document Structure', level=1)
            
            summary = layout.get('summary', {})
            type_counts = summary.get('type_counts', {})
            
            if type_counts:
                struct_table = doc.add_table(rows=len(type_counts) + 1, cols=2)
                struct_table.style = 'Table Grid'
                
                struct_table.rows[0].cells[0].text = 'Element Type'
                struct_table.rows[0].cells[1].text = 'Count'
                
                for i, (elem_type, count) in enumerate(type_counts.items(), 1):
                    struct_table.rows[i].cells[0].text = elem_type.title()
                    struct_table.rows[i].cells[1].text = str(count)
            
            doc.add_paragraph()
        
        # Extracted Text
        if text_blocks:
            doc.add_heading('Extracted Text', level=1)
            
            current_page = 0
            for block in text_blocks:
                page = block.get('page', 1)
                if page != current_page:
                    doc.add_heading(f'Page {page}', level=2)
                    current_page = page
                
                text = block.get('text', '')
                p = doc.add_paragraph(text)
                p.paragraph_format.space_after = Pt(6)
        
        # Tables
        if tables:
            doc.add_heading('Extracted Tables', level=1)
            
            for idx, table_data in enumerate(tables):
                doc.add_heading(format_table_export_title(idx + 1, table_data), level=2)
                
                meta_p = doc.add_paragraph()
                meta_p.add_run(f'{table_data.get("rows", "?")} rows × {table_data.get("columns", "?")} columns')
                
                data = table_data.get('data', [])
                if data:
                    rows = len(data)
                    cols = len(data[0]) if data else 0
                    
                    if rows > 0 and cols > 0:
                        table = doc.add_table(rows=rows, cols=cols)
                        table.style = 'Table Grid'
                        
                        for i, row_data in enumerate(data):
                            row = table.rows[i]
                            for j, cell_text in enumerate(row_data):
                                if j < len(row.cells):
                                    row.cells[j].text = str(cell_text)
                
                doc.add_paragraph()
        
        # Keywords
        if keywords:
            doc.add_heading('Keywords', level=1)
            keywords_p = doc.add_paragraph()
            keywords_p.add_run(", ".join(keywords))
        
        # Footer
        doc.add_paragraph()
        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_p.add_run("Generated by DocuVision - Intelligent Document Processing System")
        footer_run.italic = True
        footer_run.font.size = Pt(9)
        
        doc.save(output_path)
        
        logger.info(f"Exported Word document to {output_path}")
        return output_path
